#!/usr/bin/env python3
"""Membangun buku kasus uji QA (Excel) dari katalog kasus di berkas kasus_*.py.

    python3 manual/build_workbook.py            # dari folder qa/; hasil: manual/Buku-Kasus-Uji-QA.xlsx

Penguji TIDAK perlu menjalankan skrip ini; cukup membuka berkas XLSX. Skrip ini dipakai tim teknis untuk membuat ulang
buku bila kasus uji berubah. Ia juga memeriksa: (1) setiap ID kasus yang tertulis di skrip otomatis ada di katalog,
(2) setiap fitur yang tersedia/sebagian punya minimal satu kasus uji, (3) tidak ada kode fitur yang salah ketik.
Hanya butuh openpyxl (pip install openpyxl).
"""
import math
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.datavalidation import DataValidation

AKAR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))

import cases  # noqa: E402
import fitur  # noqa: E402
from kasus_dasar import URUTAN_AREA  # noqa: E402

KELUARAN = Path(__file__).resolve().parent / "Buku-Kasus-Uji-QA.xlsx"

# ---------------------------------------------------------------------------------------------------------- gaya
FONT = "Arial"
NAVY = "1F3864"
KUNING = "FFF2CC"
ABU = "EDEDED"
ABU_MUDA = "F7F7F7"
BIRU_MUDA = "DDEBF7"
HIJAU = "C6EFCE"
MERAH = "FFC7CE"
JINGGA = "FFEB9C"
ABU_TUA = "D9D9D9"

TIPIS = Side(style="thin", color="BFBFBF")
GARIS = Border(left=TIPIS, right=TIPIS, top=TIPIS, bottom=TIPIS)


def f(bold=False, italic=False, size=10, color="000000"):
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color)


def isi(warna):
    return PatternFill("solid", start_color=warna, end_color=warna)


RATA_KIRI = Alignment(horizontal="left", vertical="top", wrap_text=True)
TENGAH = Alignment(horizontal="center", vertical="center", wrap_text=True)
TENGAH_ATAS = Alignment(horizontal="center", vertical="top", wrap_text=True)


def judul(ws, teks, sub=None, lebar=8):
    ws["A1"] = teks
    ws["A1"].font = f(True, size=14, color=NAVY)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = f(italic=True, color="595959")
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=lebar)
        ws.row_dimensions[2].height = 30


def kepala(ws, baris, daftar, tinggi=32):
    for i, nama in enumerate(daftar, start=1):
        c = ws.cell(row=baris, column=i, value=nama)
        c.font = f(True, color="FFFFFF")
        c.fill = isi(NAVY)
        c.alignment = TENGAH
        c.border = GARIS
    ws.row_dimensions[baris].height = tinggi


def lebar_kolom(ws, daftar):
    for i, w in enumerate(daftar, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def tinggi_baris(teks_dan_lebar, minimum=15, per_baris=12.6):
    """Perkiraan tinggi baris (poin) dari teks yang dibungkus pada lebar kolom tertentu (satuan lebar kolom Excel)."""
    maks = 1
    for teks, lebar in teks_dan_lebar:
        if not teks:
            continue
        kolom = max(4, int(lebar * 1.05))
        n = 0
        for baris in str(teks).split("\n"):
            n += max(1, math.ceil(len(baris) / kolom))
        maks = max(maks, n)
    return min(409, max(minimum, maks * per_baris + 4))


def daftar_pilihan(ws, rentang, pilihan, pesan_salah=None):
    dv = DataValidation(type="list", formula1='"' + ",".join(pilihan) + '"', allow_blank=True)
    dv.error = pesan_salah or "Pilih salah satu dari daftar."
    dv.errorTitle = "Nilai tidak valid"
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(rentang)


def warna_status(ws, rentang, peta):
    for nilai, (latar, huruf) in peta.items():
        ws.conditional_formatting.add(
            rentang,
            CellIsRule(operator="equal", formula=[f'"{nilai}"'], fill=isi(latar), font=Font(name=FONT, size=10, color=huruf, bold=True)),
        )


STATUS_UJI = ["Lulus", "Gagal", "Diblokir", "Dilewati", "Belum diuji"]
WARNA_UJI = {"Lulus": (HIJAU, "006100"), "Gagal": (MERAH, "9C0006"), "Diblokir": (JINGGA, "7F6000"), "Dilewati": (ABU_TUA, "404040")}
STATUS_CEK = ["Lulus", "Gagal", "Diblokir", "Tidak diuji", "Tidak berlaku"]
WARNA_CEK = {"Lulus": (HIJAU, "006100"), "Gagal": (MERAH, "9C0006"), "Diblokir": (JINGGA, "7F6000"), "Tidak berlaku": (ABU_TUA, "404040")}

# ------------------------------------------------------------------------------------ pemeriksaan silang dengan skrip


def cari_id_skrip():
    """ID kasus yang tertulis di judul uji Playwright dan di skrip API, beserta berkas sumbernya."""
    peta = defaultdict(set)
    for berkas in sorted((AKAR / "e2e").rglob("*.ts")):
        isi_berkas = berkas.read_text(encoding="utf-8")
        for id_ in re.findall(r"\[(TC-[A-Z]+-\d+)\]", isi_berkas):
            peta[id_].add(str(berkas.relative_to(AKAR / "e2e")))
    smoke = AKAR / "api" / "smoke.mjs"
    if smoke.exists():
        for id_ in re.findall(r'uji\("(TC-API-\d+)"', smoke.read_text(encoding="utf-8")):
            lengkap = int(id_.split("-")[2]) >= 21
            peta[id_].add("api/smoke.mjs --lengkap" if lengkap else "api/smoke.mjs")
    return peta


SKRIP = cari_id_skrip()
ID_KATALOG = {k.id for k in cases.SEMUA}
TAK_ADA = sorted(set(SKRIP) - ID_KATALOG)
if TAK_ADA:
    sys.exit(f"ID di skrip otomatis yang tidak ada di katalog kasus: {TAK_ADA}")

KODE_FITUR = {k for k, _, _, _ in fitur.FITUR}
RUJUKAN = defaultdict(list)
for k in cases.SEMUA:
    for kode in re.findall(r"\b[A-Z]{2}-\d{2}\b", k.fitur):
        if kode not in KODE_FITUR:
            sys.exit(f"{k.id}: kode fitur tidak dikenal: {kode}")
        RUJUKAN[kode].append(k.id)
TAK_TERCAKUP = [f"{kode} {nama} [{status}]" for kode, nama, status, _ in fitur.FITUR if kode not in RUJUKAN]
if TAK_TERCAKUP:
    sys.exit("Fitur tanpa kasus uji:\n  " + "\n  ".join(TAK_TERCAKUP))

DAFTAR = cases.urut()
BARIS_AWAL = 5
BARIS_AKHIR = BARIS_AWAL + len(DAFTAR) - 1
KU = "'Kasus Uji'"


def rng(kolom):
    return f"{KU}!${kolom}${BARIS_AWAL}:${kolom}${BARIS_AKHIR}"


wb = Workbook()

# ============================================================================================ lembar: Petunjuk
ws = wb.active
ws.title = "Petunjuk"
lebar_kolom(ws, [4, 26, 70, 34, 34])
ws.sheet_view.showGridLines = False
ws["B1"] = "Buku Kasus Uji dan QC: Aplikasi Belajar Bahasa Jepang"
ws["B1"].font = f(True, size=16, color=NAVY)
ws["B2"] = f"{len(DAFTAR)} kasus uji, {len(fitur.FITUR)} fitur terpetakan. Dibuat 8 Oktober 2026 dari kode aplikasi dan hasil uji di lingkungan pengembangan; belum dijalankan di server penguji."
ws["B2"].font = f(italic=True, color="595959")
ws.merge_cells("B2:E2")

baris = 4


def seksi(teks):
    global baris
    c = ws.cell(row=baris, column=2, value=teks)
    c.font = f(True, size=11, color="FFFFFF")
    for kol in range(2, 6):
        ws.cell(row=baris, column=kol).fill = isi(NAVY)
    baris += 1


def teks_baris(label, isi_teks, tinggi=None):
    global baris
    a = ws.cell(row=baris, column=2, value=label)
    a.font = f(True)
    a.alignment = RATA_KIRI
    b = ws.cell(row=baris, column=3, value=isi_teks)
    b.font = f()
    b.alignment = RATA_KIRI
    ws.merge_cells(start_row=baris, start_column=3, end_row=baris, end_column=5)
    ws.row_dimensions[baris].height = tinggi or tinggi_baris([(isi_teks, 138)])
    baris += 1


seksi("1. Untuk apa buku ini")
teks_baris("Tujuan", "Memandu penguji memeriksa aplikasi murid, aplikasi admin, dan API secara menyeluruh sebelum rilis, mencatat hasilnya, mencatat bug, dan memutuskan layak rilis atau tidak. "
           "Kasus yang juga dijalankan skrip otomatis ditandai di kolom Otomatis (lihat qa/README.md untuk cara menjalankannya); kasus lainnya dijalankan manual.")
teks_baris("Yang belum dicakup", "Kasus uji di sini disusun dari kode dan pengujian di lingkungan pengembangan dengan layanan AI tiruan. Mutu balasan AI sungguhan, suara sungguhan, "
           "email sungguhan, Safari/Firefox, dan ponsel sungguhan baru bisa dinilai oleh penguji. Hasil di server penguji bisa berbeda: catat setiap perbedaan.")

baris += 1
seksi("2. Cara memakai (urutan yang disarankan)")
for i, langkah in enumerate([
    "Isi lembar Lingkungan (alamat, versi, akun uji). Siapkan akun uji sesuai qa/README.md.",
    "Jalankan uji otomatis (npm run test:api, lalu npm test) dan catat hasilnya di lembar Ringkasan, blok \"Hasil uji otomatis\".",
    "Kerjakan lembar Kasus Uji dari atas ke bawah, dahulukan prioritas P1. Isi kolom kuning: Hasil, Hasil aktual, ID Bug, Penguji, Tanggal, Perangkat.",
    "Setiap kasus yang Gagal: catat di lembar Log Bug (satu baris per bug) dan tulis ID bug-nya di kolom ID Bug pada kasusnya.",
    "Uji alur inti di beberapa perangkat dan browser, lalu isi lembar Perangkat.",
    "Periksa isi pelajaran bersama pengajar (kasus TC-KT-*, berkas reports/konten dari npm run konten).",
    "Sebelum rilis, kerjakan lembar QC Rilis dan isi keputusan akhir. Lembar Ringkasan menghitung sendiri angka-angkanya.",
], start=1):
    teks_baris(f"Langkah {i}", langkah)

baris += 1
seksi("3. Warna dan kolom")
for label, warna, ket in [
    ("Kuning", KUNING, "Sel yang DIISI penguji (Hasil, catatan, ID Bug, penguji, tanggal, perangkat, status, dan sebagainya)."),
    ("Abu-abu", ABU, "Sel hasil hitungan otomatis (rumus). Jangan diketik ulang."),
    ("Putih", "FFFFFF", "Isi kasus uji dari tim (judul, langkah, hasil yang diharapkan). Bila langkah atau harapannya keliru, catat di kolom Catatan penguji dan kabari tim."),
]:
    a = ws.cell(row=baris, column=2, value=label)
    a.font = f(True)
    a.fill = isi(warna)
    a.border = GARIS
    b = ws.cell(row=baris, column=3, value=ket)
    b.font = f()
    b.alignment = RATA_KIRI
    ws.merge_cells(start_row=baris, start_column=3, end_row=baris, end_column=5)
    ws.row_dimensions[baris].height = tinggi_baris([(ket, 138)])
    baris += 1

baris += 1
seksi("4. Arti status hasil kasus uji")
for label, ket in [
    ("Lulus", "Semua langkah dijalankan dan hasilnya sama dengan Hasil yang diharapkan."),
    ("Gagal", "Hasilnya berbeda dari yang diharapkan atau terjadi galat. Wajib dicatat di Log Bug."),
    ("Diblokir", "Tidak bisa dijalankan karena hal lain (mis. layanan AI belum aktif, data belum siap). Tulis penyebabnya di Hasil aktual."),
    ("Dilewati", "Sengaja tidak dijalankan pada putaran ini (mis. butuh perangkat yang tidak ada). Tulis alasannya."),
    ("Belum diuji", "Bawaan: belum dijalankan."),
]:
    teks_baris(label, ket)

baris += 1
seksi("5. Tingkat keparahan bug")
for label, ket in [
    ("Kritis", "Aplikasi tidak bisa dipakai, data hilang/bocor, atau murid tidak bisa belajar sama sekali. Menghentikan rilis."),
    ("Tinggi", "Fitur utama rusak atau hasilnya salah (mis. nilai/XP salah) dan tidak ada jalan pintas. Sebaiknya diperbaiki sebelum rilis."),
    ("Sedang", "Fitur tidak berfungsi seperti seharusnya tetapi ada jalan pintas, atau berdampak pada sedikit pengguna."),
    ("Rendah", "Kosmetik, salah ketik, atau gangguan kecil yang tidak menghalangi belajar."),
]:
    teks_baris(label, ket)

baris += 1
seksi("6. Isi lembar")
for label, ket in [
    ("Lingkungan", "Alamat, versi, pengaturan server, dan akun uji untuk putaran ini."),
    ("Kasus Uji", "Seluruh kasus uji: langkah, hasil yang diharapkan, dan kolom isian hasil."),
    ("Ringkasan", "Hitungan otomatis hasil per area dan prioritas, bug terbuka, serta rekomendasi sementara."),
    ("Log Bug", "Daftar bug yang ditemukan."),
    ("Perangkat", "Matriks alur inti per perangkat dan browser."),
    ("QC Rilis", "Daftar periksa sebelum rilis dan keputusan akhir."),
    ("Cakupan Fitur", "Pemetaan 111 fitur ke kasus uji; fitur yang tersedia tetapi tanpa kasus uji ditandai merah."),
    ("Data Uji", "Contoh isian untuk uji batas, kalkulator ambang lulus, dan jumlah soal tiap pelajaran."),
    ("Temuan Awal", "Hal yang ditemukan tim saat menyusun buku ini dan perlu keputusan. Bukan bug yang sudah dikonfirmasi di server penguji."),
]:
    teks_baris(label, ket)

baris += 1
seksi("7. Contoh pengisian (hanya contoh, bukan data nyata)")
contoh = [
    ("Kasus Uji", "ID TC-PL-10 | Hasil: Gagal | Hasil aktual: \"Layar hasil menampilkan 2 bintang padahal 0 salah; XP +20 bukan +30\" | ID Bug: BUG-003 | Penguji: Sari | Tanggal: 09/10/2026 | Perangkat: Pixel 7, Chrome 130"),
    ("Log Bug", "BUG-003 | 09/10/2026 | Sari | TC-PL-10 | Bintang salah di pelajaran 1 | Langkah: 1) Masuk akun baru, 2) kerjakan pelajaran 1 tanpa salah | Seharusnya: 3 bintang +30 XP | Aktual: 2 bintang +20 XP | Keparahan: Tinggi | Status: Baru"),
    ("Perangkat", "Baris \"iPhone, Safari 18\": kolom Pelajaran sampai hasil = Lulus; kolom Ngobrol AI (suara) = Gagal; Catatan: \"tombol mikrofon mati, pesan HTTPS muncul\""),
]
for label, ket in contoh:
    a = ws.cell(row=baris, column=2, value=label)
    a.font = f(True, italic=True, color="7F7F7F")
    b = ws.cell(row=baris, column=3, value=ket)
    b.font = f(italic=True, color="7F7F7F")
    b.alignment = RATA_KIRI
    ws.merge_cells(start_row=baris, start_column=3, end_row=baris, end_column=5)
    ws.row_dimensions[baris].height = tinggi_baris([(ket, 138)])
    baris += 1

# ============================================================================================ lembar: Lingkungan
ws = wb.create_sheet("Lingkungan")
judul(ws, "Lingkungan uji", "Isi sebelum mulai menguji. Kolom Contoh hanya memperlihatkan bentuk isian yang diharapkan.", 4)
lebar_kolom(ws, [34, 44, 40, 40])
kepala(ws, 4, ["Hal", "Isian penguji", "Contoh (bukan data nyata)", "Keterangan"])
ISIAN = [
    ("Nama lingkungan", "Staging", "Server percobaan sebelum dipakai murid."),
    ("Tanggal uji (mulai - selesai)", "09/10/2026 - 11/10/2026", ""),
    ("Penguji utama", "Sari Wulandari", ""),
    ("Alamat aplikasi murid", "https://belajar.contoh.id", "Sama dengan QA_STUDENT_URL di qa/.env."),
    ("Alamat aplikasi admin", "https://admin-belajar.contoh.id", "Sama dengan QA_ADMIN_URL."),
    ("Alamat API", "https://belajar.contoh.id/api", "Sama dengan QA_API_URL. Di server produksi API diakses lewat awalan /api pada alamat aplikasi."),
    ("Versi aplikasi (commit)", "b741aca", "Hasil perintah: git rev-parse --short HEAD di server atau repositori."),
    ("Waktu deploy terakhir", "08/10/2026 21:15", ""),
    ("Layanan AI aktif (kunci OpenAI)?", "Ya", "Ya/Tidak. Bila Tidak, kasus TC-AI-09 dan seterusnya diblokir."),
    ("Suara (TTS) sudah dibuat?", "Tidak", "Perintah tts:seed. Bila Tidak, audio pelajaran/Kamus belum ada."),
    ("Jatah AI harian (TUTOR_DAILY_QUOTA)", "20", "Turunkan (mis. 2) bila menguji jatah habis."),
    ("Zona waktu server (TZ)", "Asia/Jakarta", "Menentukan pergantian hari streak, awal minggu peringkat, dan jatah AI."),
    ("Pengiriman email aktif?", "Tidak (log saja)", "Saat ini email undangan/reset hanya dicatat di log API."),
    ("Jumlah murid / kelas sebelum uji", "8 murid / 13 kelas", "Dari lembar admin Murid dan Kelas."),
    ("Versi Node.js dan Playwright", "Node 22.11, Playwright 1.56.1", "Perintah: node -v dan npx playwright --version."),
    ("Alat lain (browser, ponsel)", "Chrome 130, Pixel 7, iPhone 14 iOS 18", "Detail perangkat dicatat juga di lembar Perangkat."),
]
for i, (hal, contoh_isi, ket) in enumerate(ISIAN, start=5):
    a = ws.cell(row=i, column=1, value=hal)
    a.font = f(True)
    a.alignment = RATA_KIRI
    a.border = GARIS
    b = ws.cell(row=i, column=2)
    b.fill = isi(KUNING)
    b.border = GARIS
    b.font = f()
    b.alignment = RATA_KIRI
    c = ws.cell(row=i, column=3, value=contoh_isi)
    c.font = f(italic=True, color="7F7F7F")
    c.alignment = RATA_KIRI
    c.border = GARIS
    d = ws.cell(row=i, column=4, value=ket)
    d.font = f(color="595959")
    d.alignment = RATA_KIRI
    d.border = GARIS
    ws.row_dimensions[i].height = tinggi_baris([(ket, 40), (hal, 34)])

akhir = 5 + len(ISIAN) + 1
ws.cell(row=akhir, column=1, value="Akun uji (JANGAN menulis password di berkas ini)").font = f(True, size=11, color=NAVY)
kepala(ws, akhir + 1, ["Peran", "Email", "Contoh", "Dipakai untuk"], 22)
AKUN = [
    ("Murid uji utama (QA_STUDENT_*)", "qa-murid@example.com", "Semua uji murid dan uji otomatis; progres belajar akan terisi."),
    ("Murid uji kedua (QA_STUDENT2_*)", "qa-murid2@example.com", "Uji yang mengubah akun: pindah kelas, nonaktifkan, reset password."),
    ("Murid uji penguncian (QA_LOCKOUT_*)", "qa-murid-kunci@example.com", "HANYA uji penguncian akun; terkunci 15 menit sesudahnya."),
    ("Admin uji (QA_ADMIN_*)", "qa-admin@example.com", "Semua uji admin."),
    ("Murid baru (dari undangan)", "qa-baru-1@example.com", "Uji pendaftaran lewat undangan (TC-AK-01 dan seterusnya)."),
]
for i, (peran, contoh_email, guna) in enumerate(AKUN, start=akhir + 2):
    ws.cell(row=i, column=1, value=peran).font = f(True)
    kuning = ws.cell(row=i, column=2)
    kuning.fill = isi(KUNING)
    kuning.border = GARIS
    c = ws.cell(row=i, column=3, value=contoh_email)
    c.font = f(italic=True, color="7F7F7F")
    ws.cell(row=i, column=4, value=guna).font = f(color="595959")
    for kol in (1, 3, 4):
        ws.cell(row=i, column=kol).alignment = RATA_KIRI
        ws.cell(row=i, column=kol).border = GARIS
    ws.row_dimensions[i].height = tinggi_baris([(guna, 40), (peran, 34)])
ws.freeze_panes = "A5"

# ================================================================================================ lembar: Kasus Uji
ws = wb.create_sheet("Kasus Uji")
KOLOM_KU = [
    ("No", 5), ("ID", 11), ("Area", 18), ("Aplikasi", 9), ("Prioritas", 9), ("Judul kasus uji", 38), ("Prasyarat", 32), ("Langkah", 52),
    ("Hasil yang diharapkan", 52), ("Fitur terkait", 15), ("Otomatis", 26), ("Catatan dari tim", 34),
    ("Hasil", 13), ("Hasil aktual / catatan penguji", 36), ("ID Bug", 11), ("Penguji", 12), ("Tanggal uji", 12), ("Perangkat / browser", 20),
]
judul(ws, "Kasus uji", "Kolom kuning diisi penguji. Prioritas: P1 = kritis (alur inti dan keamanan), P2 = penting, P3 = pelengkap. "
      "Kolom Otomatis menunjukkan skrip yang menjalankan kasus yang sama; tetap isi Hasil sesuai laporan skrip atau hasil uji manual Anda.", 12)
lebar_kolom(ws, [w for _, w in KOLOM_KU])
kepala(ws, 4, [n for n, _ in KOLOM_KU], 34)
KOL_HASIL = 13
for i, k in enumerate(DAFTAR):
    r = BARIS_AWAL + i
    langkah = "\n".join(f"{n}. {t}" for n, t in enumerate(k.langkah, start=1))
    auto = "Ya: " + "; ".join(sorted(SKRIP[k.id])) if k.id in SKRIP else "Tidak (manual)"
    nilai = [i + 1, k.id, k.area, k.app, k.prio, k.judul, k.pra, langkah, k.harap, k.fitur, auto, k.catatan]
    for kol, v in enumerate(nilai, start=1):
        c = ws.cell(row=r, column=kol, value=v)
        c.font = f(bold=(kol == 2))
        c.alignment = TENGAH_ATAS if kol in (1, 2, 4, 5) else RATA_KIRI
        c.border = GARIS
        if kol == 11 and v == "Tidak (manual)":
            c.font = f(color="7F7F7F")
    h = ws.cell(row=r, column=KOL_HASIL, value="Belum diuji")
    for kol in range(KOL_HASIL, len(KOLOM_KU) + 1):
        c = ws.cell(row=r, column=kol)
        c.fill = isi(KUNING)
        c.border = GARIS
        c.font = f(bold=(kol == KOL_HASIL))
        c.alignment = TENGAH_ATAS if kol in (KOL_HASIL, 15, 16, 17) else RATA_KIRI
    ws.cell(row=r, column=17).number_format = "dd/mm/yyyy"
    ws.row_dimensions[r].height = tinggi_baris([(k.judul, 38), (k.pra, 32), (langkah, 52), (k.harap, 52), (k.catatan, 34), (auto, 26)])

daftar_pilihan(ws, f"M{BARIS_AWAL}:M{BARIS_AKHIR}", STATUS_UJI)
warna_status(ws, f"M{BARIS_AWAL}:M{BARIS_AKHIR}", WARNA_UJI)
for prio, warna in (("P1", "F8CBAD"), ("P2", "FFE699"), ("P3", "E2EFDA")):
    ws.conditional_formatting.add(f"E{BARIS_AWAL}:E{BARIS_AKHIR}", CellIsRule(operator="equal", formula=[f'"{prio}"'], fill=isi(warna)))
ws.freeze_panes = f"C{BARIS_AWAL}"
ws.auto_filter.ref = f"A4:{get_column_letter(len(KOLOM_KU))}{BARIS_AKHIR}"
ws.print_title_rows = "4:4"

# ================================================================================================== lembar: Log Bug
ws = wb.create_sheet("Log Bug")
KOLOM_BUG = [
    ("ID Bug", 10), ("Tanggal", 12), ("Pelapor", 14), ("ID kasus uji", 12), ("Aplikasi", 11), ("Judul singkat", 34), ("Langkah mereproduksi", 48),
    ("Hasil yang seharusnya", 34), ("Hasil aktual", 34), ("Keparahan", 11), ("Prioritas perbaikan", 15), ("Perangkat / browser", 20), ("Lingkungan / versi", 20),
    ("Lampiran (nama berkas)", 24), ("Status", 16), ("Penanggung jawab", 16), ("Catatan tim teknis", 30),
]
JUMLAH_BUG = 150
BUG_AWAL, BUG_AKHIR = 5, 5 + JUMLAH_BUG - 1
judul(ws, "Log bug", "Satu baris per bug. Isi sebanyak mungkin kolom kuning; sertakan tangkapan layar (nama berkasnya di kolom Lampiran). Contoh pengisian ada di lembar Petunjuk.", 10)
lebar_kolom(ws, [w for _, w in KOLOM_BUG])
kepala(ws, 4, [n for n, _ in KOLOM_BUG], 34)
for i in range(JUMLAH_BUG):
    r = BUG_AWAL + i
    for kol in range(1, len(KOLOM_BUG) + 1):
        c = ws.cell(row=r, column=kol)
        c.border = GARIS
        c.font = f(bold=(kol == 1))
        c.alignment = TENGAH_ATAS if kol in (1, 2, 4, 5, 10, 11) else RATA_KIRI
        if kol != 1:
            c.fill = isi(KUNING)
    ws.cell(row=r, column=1, value=f"BUG-{i + 1:03d}")
    ws.cell(row=r, column=2).number_format = "dd/mm/yyyy"
daftar_pilihan(ws, f"E{BUG_AWAL}:E{BUG_AKHIR}", ["Murid", "Admin", "API", "Konten", "Semua"])
daftar_pilihan(ws, f"J{BUG_AWAL}:J{BUG_AKHIR}", ["Kritis", "Tinggi", "Sedang", "Rendah"])
daftar_pilihan(ws, f"K{BUG_AWAL}:K{BUG_AKHIR}", ["Segera", "Sebelum rilis", "Setelah rilis"])
daftar_pilihan(ws, f"O{BUG_AWAL}:O{BUG_AKHIR}", ["Baru", "Dikonfirmasi", "Sedang diperbaiki", "Siap diuji ulang", "Ditutup", "Ditolak", "Ditunda"])
warna_status(ws, f"J{BUG_AWAL}:J{BUG_AKHIR}", {"Kritis": (MERAH, "9C0006"), "Tinggi": ("F8CBAD", "833C0B"), "Sedang": (JINGGA, "7F6000"), "Rendah": (HIJAU, "006100")})
warna_status(ws, f"O{BUG_AWAL}:O{BUG_AKHIR}", {"Ditutup": (HIJAU, "006100"), "Ditolak": (ABU_TUA, "404040"), "Ditunda": (ABU_TUA, "404040")})
ws.freeze_panes = f"C{BUG_AWAL}"
ws.auto_filter.ref = f"A4:{get_column_letter(len(KOLOM_BUG))}{BUG_AKHIR}"
LB = "'Log Bug'"

# =============================================================================================== lembar: Perangkat
ws = wb.create_sheet("Perangkat")
ALUR = ["Masuk dan keluar", "Beranda", "Pelajaran sampai hasil", "Percakapan (skenario)", "Ngobrol AI (ketik)", "Ngobrol AI (suara)", "Kamus", "Flashcard", "Leaderboard", "Profil", "Admin"]
judul(ws, "Perangkat dan browser", "Uji alur inti di tiap perangkat dan browser yang dipakai murid. Skrip otomatis hanya mencakup Chromium (Chrome/Edge) dan emulasi ponsel; "
      "Safari, Firefox, dan ponsel sungguhan wajib diuji manual. Kolom Admin hanya untuk perangkat komputer.", 8)
KOL_PERANGKAT = ["Perangkat / sistem operasi", "Browser dan versi"] + ALUR + ["Catatan", "Penguji", "Tanggal"]
lebar_kolom(ws, [28, 22] + [11] * len(ALUR) + [36, 12, 12])
kepala(ws, 4, KOL_PERANGKAT, 46)
PRESET = [
    ("Windows 11", "Chrome (terbaru)"), ("Windows 11", "Edge (terbaru)"), ("Windows 11", "Firefox (terbaru)"), ("macOS", "Safari (terbaru)"), ("macOS", "Chrome (terbaru)"),
    ("Android 13 atau lebih baru (ponsel)", "Chrome"), ("Android (ponsel, layar kecil < 360 px)", "Chrome"), ("Android", "Samsung Internet"),
    ("iPhone (iOS 17 atau lebih baru)", "Safari"), ("iPhone", "Chrome"), ("iPad", "Safari"),
]
JUMLAH_PERANGKAT = len(PRESET) + 5
P_AWAL, P_AKHIR = 5, 5 + JUMLAH_PERANGKAT - 1
for i in range(JUMLAH_PERANGKAT):
    r = P_AWAL + i
    nilai_awal = PRESET[i] if i < len(PRESET) else ("", "")
    for kol in range(1, len(KOL_PERANGKAT) + 1):
        c = ws.cell(row=r, column=kol)
        c.border = GARIS
        c.font = f()
        c.alignment = TENGAH_ATAS if 3 <= kol <= 2 + len(ALUR) else RATA_KIRI
        c.fill = isi(KUNING)
    ws.cell(row=r, column=1, value=nilai_awal[0] or None)
    ws.cell(row=r, column=2, value=nilai_awal[1] or None)
    for kol in range(3, 3 + len(ALUR)):
        ws.cell(row=r, column=kol, value="Tidak diuji")
    ws.cell(row=r, column=len(KOL_PERANGKAT) - 0).number_format = "dd/mm/yyyy"
    ws.row_dimensions[r].height = 30
alur_rng = f"C{P_AWAL}:{get_column_letter(2 + len(ALUR))}{P_AKHIR}"
daftar_pilihan(ws, alur_rng, STATUS_CEK)
warna_status(ws, alur_rng, WARNA_CEK)
r_sum = P_AKHIR + 2
ws.cell(row=r_sum, column=1, value="Jumlah Lulus").font = f(True)
ws.cell(row=r_sum + 1, column=1, value="Jumlah Gagal").font = f(True)
for kol in range(3, 3 + len(ALUR)):
    huruf = get_column_letter(kol)
    for off, teks in ((0, "Lulus"), (1, "Gagal")):
        c = ws.cell(row=r_sum + off, column=kol, value=f'=COUNTIF({huruf}{P_AWAL}:{huruf}{P_AKHIR},"{teks}")')
        c.font = f(True)
        c.fill = isi(ABU)
        c.alignment = TENGAH
        c.border = GARIS
ws.cell(row=r_sum + 3, column=1, value="Catatan: perangkat tanpa fitur tertentu (mis. komputer tanpa mikrofon untuk Ngobrol AI suara) diisi \"Tidak berlaku\".").font = f(italic=True, color="595959")
ws.freeze_panes = "C5"

# =============================================================================================== lembar: QC Rilis
ws = wb.create_sheet("QC Rilis")
judul(ws, "QC sebelum rilis", "Item bertanda Wajib = Ya harus Lulus (atau Tidak berlaku dengan alasan) sebelum rilis. Dua item pertama di bagian Pengujian terisi otomatis dari lembar Ringkasan.", 9)
lebar_kolom(ws, [5, 22, 60, 8, 34, 15, 18, 12, 36])
KOL_QC = ["No", "Kelompok", "Pemeriksaan", "Wajib", "Cara memeriksa / rujukan kasus", "Status", "Penanggung jawab", "Tanggal", "Catatan"]
QC = [
    ("Lingkungan dan rilis", "Versi (commit) yang diuji tercatat di lembar Lingkungan dan sama dengan yang akan dirilis", "Ya", "git rev-parse --short HEAD"),
    ("Lingkungan dan rilis", "Migrasi database berhasil dijalankan tanpa galat dan tidak ada perubahan skema yang belum dimigrasi", "Ya", "docs/DEPLOY.md; CI memeriksa migrasi"),
    ("Lingkungan dan rilis", "Seluruh variabel lingkungan wajib terisi dan rahasia kuat (JWT, database, CORS_ORIGIN_*), tidak memakai contoh bawaan", "Ya", "docs/DEPLOY.md bagian konfigurasi"),
    ("Lingkungan dan rilis", "HTTPS aktif, HTTP dialihkan, sertifikat sah", "Ya", "TC-NF-01"),
    ("Lingkungan dan rilis", "Jam server benar dan zona waktu (TZ) sesuai kebutuhan", "Ya", "TC-API-05, TC-NF-15"),
    ("Lingkungan dan rilis", "Cadangan database terbaru ada dan pernah berhasil dipulihkan di lingkungan uji", "Ya", "TC-NF-14"),
    ("Lingkungan dan rilis", "Alamat di tautan undangan dan reset password menunjuk ke aplikasi murid yang benar", "Ya", "TC-AD-30"),
    ("Pengujian", "Semua kasus uji prioritas P1 sudah dijalankan dan lulus (tidak ada Gagal atau Belum diuji)", "Ya", "Lembar Ringkasan (otomatis)"),
    ("Pengujian", "Tidak ada bug Kritis atau Tinggi yang masih terbuka", "Ya", "Lembar Ringkasan (otomatis)"),
    ("Pengujian", "Uji asap API lulus penuh (npm run test:api, dan --lengkap di lingkungan uji)", "Ya", "TC-API-01 sampai TC-API-26"),
    ("Pengujian", "Uji browser otomatis lulus tanpa kegagalan baru (npm test)", "Ya", "qa/README.md"),
    ("Pengujian", "Pemindaian aksesibilitas otomatis tidak melaporkan pelanggaran serius/kritis baru", "Ya", "TC-AX-01 sampai TC-AX-03"),
    ("Alur inti manual", "Undangan dibuat admin, murid mendaftar, lulus pelajaran pertama, XP dan Leaderboard benar", "Ya", "TC-AK-01, TC-PL-10, TC-NL-02, TC-AD-14"),
    ("Alur inti manual", "Lupa password sampai bisa masuk dengan password baru", "Ya", "TC-AK-18"),
    ("Alur inti manual", "Percakapan skenario mode Latihan dan Tes, termasuk XP satu kali", "Ya", "TC-PC-04 sampai TC-PC-08"),
    ("Alur inti manual", "Ngobrol dengan AI: kirim pesan, jatah berkurang, pesan galat jelas bila AI mati atau jatah habis", "Tidak", "TC-AI-09, TC-AI-14, TC-AI-15 (hanya bila AI dipakai di rilis ini)"),
    ("Alur inti manual", "Admin: kelas, undangan, murid, pindah kelas, nonaktifkan/aktifkan", "Ya", "TC-AD-10 sampai TC-AD-28"),
    ("Keamanan", "Pemisahan hak akses murid/admin, field terlarang ditolak, data pribadi tidak bocor", "Ya", "TC-API-19, TC-NF-06 sampai TC-NF-08"),
    ("Keamanan", "Dokumentasi API tidak terbuka ke publik; CORS hanya untuk asal resmi; header keamanan terpasang", "Ya", "TC-NF-02 sampai TC-NF-04"),
    ("Keamanan", "Pembatas permintaan aktif dan pengguna di jaringan bersama tidak terkunci sia-sia", "Ya", "TC-NF-05"),
    ("Keamanan", "Keputusan atas temuan awal: Keluar tidak mencabut sesi, murid nonaktif tetap punya sesi, XP Tes pada percobaan gagal", "Ya", "Lembar Temuan Awal; TC-AK-24, TC-AD-27, TC-PC-08"),
    ("Email", "Keputusan pengiriman email undangan dan reset password (disambungkan, atau tautan disalin dari log dengan prosedur tertulis)", "Ya", "TC-AK-31"),
    ("Konten", "Pengajar sudah meninjau isi Hiragana, Unit 3, dan skenario (berkas reports/konten); temuan konten tercatat", "Ya", "TC-KT-01 sampai TC-KT-04"),
    ("Konten", "npm run konten tidak melaporkan temuan tingkat Galat", "Ya", "TC-KT-01"),
    ("Konten", "Audio sudah dibuat dan diperiksa (bila suara dipakai di rilis ini)", "Tidak", "TC-KT-07"),
    ("Perangkat", "Alur inti lulus di Chrome desktop, Android Chrome, dan iPhone Safari", "Ya", "Lembar Perangkat; TC-RS-08, TC-RS-09"),
    ("Perangkat", "Tampilan ponsel dan layar sempit diperiksa tanpa gulir ke samping", "Tidak", "TC-RS-01 sampai TC-RS-11"),
    ("Kinerja", "Beranda termuat wajar di jaringan seluler dan beban ringan tidak menimbulkan galat", "Tidak", "TC-NF-10, TC-NF-11"),
    ("Dokumentasi", "Catatan rilis, daftar bug terbuka yang diketahui, dan perubahan penting sudah ditulis dan dibagikan", "Ya", ""),
    ("Dokumentasi", "Panduan pengguna/admin dan docs/DEPLOY.md sesuai dengan versi yang dirilis", "Tidak", "docs/DEPLOY.md, README.md"),
]
BARIS_QC = 12
kepala(ws, BARIS_QC - 1, KOL_QC, 30)
for i, (kel, periksa, wajib, rujuk) in enumerate(QC):
    r = BARIS_QC + i
    vals = [i + 1, kel, periksa, wajib, rujuk, "Belum", None, None, None]
    for kol, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=kol, value=v)
        c.font = f()
        c.border = GARIS
        c.alignment = TENGAH_ATAS if kol in (1, 4, 6, 8) else RATA_KIRI
        if kol >= 6:
            c.fill = isi(KUNING)
    ws.cell(row=r, column=8).number_format = "dd/mm/yyyy"
    ws.row_dimensions[r].height = tinggi_baris([(periksa, 60), (rujuk, 34), (kel, 22)])
QC_AKHIR = BARIS_QC + len(QC) - 1
# Dua item pengujian pertama terisi otomatis dari Ringkasan (rumus ditulis setelah lembar Ringkasan dibuat).
daftar_pilihan(ws, f"F{BARIS_QC}:F{QC_AKHIR}", ["Lulus", "Gagal", "Belum", "Tidak berlaku"])
warna_status(ws, f"F{BARIS_QC}:F{QC_AKHIR}", {"Lulus": (HIJAU, "006100"), "Gagal": (MERAH, "9C0006"), "Tidak berlaku": (ABU_TUA, "404040")})
daftar_pilihan(ws, f"D{BARIS_QC}:D{QC_AKHIR}", ["Ya", "Tidak"])
ws.freeze_panes = f"A{BARIS_QC}"

# ================================================================================================ lembar: Ringkasan
rs = wb.create_sheet("Ringkasan", 3)
judul(rs, "Ringkasan hasil", "Semua angka di lembar ini dihitung otomatis dari lembar Kasus Uji dan Log Bug. Satu-satunya isian manual: blok \"Hasil uji otomatis\".", 9)
lebar_kolom(rs, [34, 12, 10, 10, 10, 10, 12, 12, 12])

KOL_RING = ["", "Total kasus", "Lulus", "Gagal", "Diblokir", "Dilewati", "Belum diuji", "% sudah diuji", "% lulus"]


def tabel_ringkasan(awal, judul_blok, label_kolom, daftar_label, kolom_kriteria):
    rs.cell(row=awal, column=1, value=judul_blok).font = f(True, size=11, color=NAVY)
    kolom = list(KOL_RING)
    kolom[0] = label_kolom
    kepala(rs, awal + 1, kolom, 30)
    r = awal + 2
    mulai = r
    for label in daftar_label:
        rs.cell(row=r, column=1, value=label)
        k = rng(kolom_kriteria)
        rs.cell(row=r, column=2, value=f"=COUNTIFS({k},$A{r})")
        for kol, st in ((3, "Lulus"), (4, "Gagal"), (5, "Diblokir"), (6, "Dilewati")):
            rs.cell(row=r, column=kol, value=f'=COUNTIFS({k},$A{r},{rng("M")},"{st}")')
        rs.cell(row=r, column=7, value=f"=B{r}-C{r}-D{r}-E{r}-F{r}")
        rs.cell(row=r, column=8, value=f"=IF(B{r}=0,0,(C{r}+D{r})/B{r})")
        rs.cell(row=r, column=9, value=f"=IF((C{r}+D{r})=0,0,C{r}/(C{r}+D{r}))")
        r += 1
    rs.cell(row=r, column=1, value="Jumlah")
    for kol in range(2, 8):
        huruf = get_column_letter(kol)
        rs.cell(row=r, column=kol, value=f"=SUM({huruf}{mulai}:{huruf}{r - 1})")
    rs.cell(row=r, column=8, value=f"=IF(B{r}=0,0,(C{r}+D{r})/B{r})")
    rs.cell(row=r, column=9, value=f"=IF((C{r}+D{r})=0,0,C{r}/(C{r}+D{r}))")
    for rr in range(mulai, r + 1):
        for kol in range(1, 10):
            c = rs.cell(row=rr, column=kol)
            c.font = f(bold=(rr == r or kol == 1))
            c.border = GARIS
            c.alignment = Alignment(horizontal="left" if kol == 1 else "center", vertical="center", wrap_text=True)
            if kol > 1:
                c.fill = isi(ABU)
            if kol in (8, 9):
                c.number_format = "0%"
    return mulai, r


m1, t1 = tabel_ringkasan(4, "Per area", "Area", URUTAN_AREA, "C")
m2, t2 = tabel_ringkasan(t1 + 3, "Per prioritas", "Prioritas", ["P1", "P2", "P3"], "E")
# Penanda keterangan prioritas
for off, ket in enumerate(["P1 = kritis", "P2 = penting", "P3 = pelengkap"]):
    rs.cell(row=m2 + off, column=10, value=ket).font = f(italic=True, color="7F7F7F")
rs.column_dimensions["J"].width = 16

p1_baris = m2  # baris P1
b_ind = t2 + 3
rs.cell(row=b_ind, column=1, value="Indikator rilis").font = f(True, size=11, color=NAVY)
kepala(rs, b_ind + 1, ["Indikator", "Nilai"], 22)
BUG_SEV = f"{LB}!$J${BUG_AWAL}:$J${BUG_AKHIR}"
BUG_STATUS = f"{LB}!$O${BUG_AWAL}:$O${BUG_AKHIR}"
indikator = [
    ("Kasus P1 yang Gagal", f"=D{p1_baris}"),
    ("Kasus P1 yang belum selesai (Belum diuji, Diblokir, Dilewati)", f"=B{p1_baris}-C{p1_baris}-D{p1_baris}"),
    ("Bug Kritis terbuka", f'=COUNTIFS({BUG_SEV},"Kritis",{BUG_STATUS},"<>Ditutup",{BUG_STATUS},"<>Ditolak",{BUG_STATUS},"<>Ditunda")'),
    ("Bug Tinggi terbuka", f'=COUNTIFS({BUG_SEV},"Tinggi",{BUG_STATUS},"<>Ditutup",{BUG_STATUS},"<>Ditolak",{BUG_STATUS},"<>Ditunda")'),
]
ind_baris = {}
for off, (label, rumus) in enumerate(indikator):
    r = b_ind + 2 + off
    rs.cell(row=r, column=1, value=label).alignment = RATA_KIRI
    rs.cell(row=r, column=2, value=rumus)
    ind_baris[label] = r
    for kol in (1, 2):
        c = rs.cell(row=r, column=kol)
        c.border = GARIS
        c.font = f(bold=(kol == 1))
    rs.cell(row=r, column=2).fill = isi(ABU)
    rs.cell(row=r, column=2).alignment = TENGAH
    rs.row_dimensions[r].height = 28
r_rek = b_ind + 2 + len(indikator)
rs.cell(row=r_rek, column=1, value="Rekomendasi sementara (bukan keputusan)").font = f(True)
rs.cell(row=r_rek, column=1).border = GARIS
rs.cell(row=r_rek, column=1).alignment = RATA_KIRI
IND = list(ind_baris.values())
rs.cell(
    row=r_rek,
    column=2,
    value=(
        f'=IF(B{IND[0]}>0,"BELUM LAYAK: ada kasus P1 yang gagal",IF(B{IND[2]}>0,"BELUM LAYAK: ada bug Kritis terbuka",'
        f'IF(B{IND[1]}>0,"PENGUJIAN BELUM LENGKAP: kasus P1 belum selesai",IF(B{IND[3]}>0,"LAYAK DENGAN CATATAN: masih ada bug Tinggi terbuka","LAYAK secara pengujian: lanjutkan ke lembar QC Rilis"))))'
    ),
)
rs.merge_cells(start_row=r_rek, start_column=2, end_row=r_rek, end_column=9)
rs.cell(row=r_rek, column=2).font = f(True)
rs.cell(row=r_rek, column=2).fill = isi(ABU)
rs.cell(row=r_rek, column=2).alignment = Alignment(wrap_text=True, vertical="center")
rs.cell(row=r_rek, column=2).border = GARIS
rs.row_dimensions[r_rek].height = 30
rs.conditional_formatting.add(f"B{r_rek}", FormulaRule(formula=[f'LEFT(B{r_rek},5)="BELUM"'], fill=isi(MERAH)))
rs.conditional_formatting.add(f"B{r_rek}", FormulaRule(formula=[f'LEFT(B{r_rek},5)="LAYAK"'], fill=isi(HIJAU)))
rs.conditional_formatting.add(f"B{r_rek}", FormulaRule(formula=[f'LEFT(B{r_rek},9)="PENGUJIAN"'], fill=isi(JINGGA)))
# Bug dari Log Bug
b_bug = r_rek + 3
rs.cell(row=b_bug, column=1, value="Bug menurut keparahan").font = f(True, size=11, color=NAVY)
kepala(rs, b_bug + 1, ["Keparahan", "Total", "Terbuka", "Selesai atau ditolak"], 30)
for off, sev in enumerate(["Kritis", "Tinggi", "Sedang", "Rendah"]):
    r = b_bug + 2 + off
    rs.cell(row=r, column=1, value=sev)
    rs.cell(row=r, column=2, value=f'=COUNTIFS({BUG_SEV},$A{r})')
    rs.cell(row=r, column=3, value=f'=COUNTIFS({BUG_SEV},$A{r},{BUG_STATUS},"<>Ditutup",{BUG_STATUS},"<>Ditolak",{BUG_STATUS},"<>Ditunda")')
    rs.cell(row=r, column=4, value=f"=B{r}-C{r}")
    for kol in range(1, 5):
        c = rs.cell(row=r, column=kol)
        c.border = GARIS
        c.font = f(bold=(kol == 1))
        c.alignment = Alignment(horizontal="left" if kol == 1 else "center", vertical="center")
        if kol > 1:
            c.fill = isi(ABU)
r_tot_bug = b_bug + 6
rs.cell(row=r_tot_bug, column=1, value="Jumlah")
for kol, huruf in ((2, "B"), (3, "C"), (4, "D")):
    rs.cell(row=r_tot_bug, column=kol, value=f"=SUM({huruf}{b_bug + 2}:{huruf}{b_bug + 5})")
for kol in range(1, 5):
    c = rs.cell(row=r_tot_bug, column=kol)
    c.font = f(True)
    c.border = GARIS
    c.alignment = Alignment(horizontal="left" if kol == 1 else "center")
    if kol > 1:
        c.fill = isi(ABU)
# Hasil uji otomatis (isian manual)
b_auto = r_tot_bug + 3
rs.cell(row=b_auto, column=1, value="Hasil uji otomatis (isi dari laporan skrip)").font = f(True, size=11, color=NAVY)
kepala(rs, b_auto + 1, ["Skrip", "Total", "Lulus", "Gagal", "Dilewati", "Tanggal", "% lulus"], 30)
for off, (nama, contoh_baris) in enumerate([("Uji browser (npm test)", None), ("Uji asap API (npm run test:api)", None), ("Uji asap API lengkap (--lengkap)", None)]):
    r = b_auto + 2 + off
    rs.cell(row=r, column=1, value=nama)
    for kol in range(2, 7):
        c = rs.cell(row=r, column=kol)
        c.fill = isi(KUNING)
    rs.cell(row=r, column=6).number_format = "dd/mm/yyyy"
    rs.cell(row=r, column=7, value=f"=IF(B{r}=0,0,C{r}/B{r})")
    rs.cell(row=r, column=7).number_format = "0%"
    rs.cell(row=r, column=7).fill = isi(ABU)
    for kol in range(1, 8):
        c = rs.cell(row=r, column=kol)
        c.border = GARIS
        c.font = f(bold=(kol == 1))
        c.alignment = Alignment(horizontal="left" if kol == 1 else "center", vertical="center", wrap_text=True)
    rs.row_dimensions[r].height = 28
rs.cell(row=b_auto + 5, column=1, value="Contoh isian: Uji browser | Total 134 | Lulus 132 | Gagal 0 | Dilewati 2 | 08/10/2026 (angka dari lingkungan pengembangan, bukan dari server Anda).").font = f(italic=True, color="7F7F7F")
rs.freeze_panes = "A4"

# Kaitkan dua item QC Rilis ke Ringkasan (rumus), sekarang setelah baris Ringkasan diketahui.
ws = wb["QC Rilis"]
r1, r2 = BARIS_QC + 7, BARIS_QC + 8
assert "P1" in ws.cell(row=r1, column=3).value and "Kritis atau Tinggi" in ws.cell(row=r2, column=3).value
ws.cell(row=r1, column=6, value=f'=IF(AND(Ringkasan!B{IND[0]}=0,Ringkasan!B{IND[1]}=0),"Lulus",IF(Ringkasan!B{IND[0]}>0,"Gagal","Belum"))')
ws.cell(row=r2, column=6, value=f'=IF(AND(Ringkasan!B{IND[2]}=0,Ringkasan!B{IND[3]}=0),"Lulus","Gagal")')
for rr in (r1, r2):
    ws.cell(row=rr, column=6).fill = isi(ABU)
    ws.cell(row=rr, column=7, value="Otomatis")
# Blok keputusan di atas daftar
ws["A3"] = "Rilis / versi"
ws["A4"] = "Tanggal keputusan"
ws["A5"] = "Penguji utama"
ws["A6"] = "Penyetuju"
ws["A7"] = "Keputusan akhir"
for r in range(3, 8):
    ws.cell(row=r, column=1).font = f(True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    c = ws.cell(row=r, column=3)
    c.fill = isi(KUNING)
    c.border = GARIS
    c.font = f()
ws["C4"].number_format = "dd/mm/yyyy"
daftar_pilihan(ws, "C7", ["Layak rilis", "Layak dengan catatan", "Tidak layak"])
ws["E3"] = "Item wajib yang belum Lulus"
ws["E3"].font = f(True)
ws["F3"] = f'=COUNTIFS(D{BARIS_QC}:D{QC_AKHIR},"Ya",F{BARIS_QC}:F{QC_AKHIR},"<>Lulus",F{BARIS_QC}:F{QC_AKHIR},"<>Tidak berlaku")'
ws["E4"] = "Item Gagal"
ws["E4"].font = f(True)
ws["F4"] = f'=COUNTIF(F{BARIS_QC}:F{QC_AKHIR},"Gagal")'
ws["E5"] = "Rekomendasi sementara"
ws["E5"].font = f(True)
ws["F5"] = f'=IF(F4>0,"BELUM LAYAK: ada item Gagal",IF(F3>0,"BELUM LENGKAP: "&F3&" item wajib belum Lulus","Semua item wajib Lulus"))'
ws.merge_cells("F5:I5")
for ref in ("F3", "F4", "F5"):
    ws[ref].font = f(True)
    ws[ref].fill = isi(ABU)
    ws[ref].border = GARIS
    ws[ref].alignment = Alignment(horizontal="center" if ref != "F5" else "left", vertical="center")
ws["E7"] = "Keputusan akhir tetap di tangan manusia; rumus di atas hanya membantu."
ws["E7"].font = f(italic=True, color="595959")
ws.merge_cells("E7:I7")

# ============================================================================================ lembar: Cakupan Fitur
ws = wb.create_sheet("Cakupan Fitur")
judul(ws, "Cakupan fitur oleh kasus uji", "Dipetakan dari dokumen Daftar Fitur Rinci (6 Oktober 2026). Fitur berstatus Belum dibangun hanya diperiksa agar tidak tampil/menjanjikan (TC-BT-01).", 6)
lebar_kolom(ws, [9, 46, 36, 11, 11, 60])
kepala(ws, 4, ["Kode", "Fitur", "Kelompok", "Status", "Jumlah kasus", "Kasus uji terkait"], 30)
F_AWAL = 5
F_AKHIR = F_AWAL + len(fitur.FITUR) - 1
for i, (kode, nama, status, kel) in enumerate(fitur.FITUR):
    r = F_AWAL + i
    ws.cell(row=r, column=1, value=kode)
    ws.cell(row=r, column=2, value=nama)
    ws.cell(row=r, column=3, value=kel)
    ws.cell(row=r, column=4, value=status)
    ws.cell(row=r, column=5, value=f'=COUNTIF({rng("J")},"*"&$A{r}&"*")')
    ws.cell(row=r, column=6, value=", ".join(RUJUKAN[kode]))
    for kol in range(1, 7):
        c = ws.cell(row=r, column=kol)
        c.font = f(bold=(kol == 1))
        c.border = GARIS
        c.alignment = TENGAH_ATAS if kol in (1, 4, 5) else RATA_KIRI
    ws.cell(row=r, column=5).fill = isi(ABU)
    ws.row_dimensions[r].height = tinggi_baris([(nama, 46), (", ".join(RUJUKAN[kode]), 60), (kel, 36)], minimum=16)
ws.conditional_formatting.add(
    f"A{F_AWAL}:F{F_AKHIR}",
    FormulaRule(formula=[f'AND($D{F_AWAL}<>"Belum",$E{F_AWAL}=0)'], fill=isi(MERAH)),
)
warna_status(ws, f"D{F_AWAL}:D{F_AKHIR}", {"Tersedia": (HIJAU, "006100"), "Sebagian": (JINGGA, "7F6000"), "Belum": (ABU_TUA, "404040")})
r_ring = F_AKHIR + 2
for off, (label, rumus) in enumerate([
    ("Jumlah fitur", f"=COUNTA(A{F_AWAL}:A{F_AKHIR})"),
    ("Tersedia", f'=COUNTIF(D{F_AWAL}:D{F_AKHIR},"Tersedia")'),
    ("Sebagian", f'=COUNTIF(D{F_AWAL}:D{F_AKHIR},"Sebagian")'),
    ("Belum dibangun", f'=COUNTIF(D{F_AWAL}:D{F_AKHIR},"Belum")'),
    ("Tersedia atau sebagian TANPA kasus uji (harus 0)", f'=COUNTIFS(D{F_AWAL}:D{F_AKHIR},"<>Belum",E{F_AWAL}:E{F_AKHIR},0)'),
]):
    a = ws.cell(row=r_ring + off, column=2, value=label)
    a.font = f(True)
    b = ws.cell(row=r_ring + off, column=5, value=rumus)
    b.font = f(True)
    b.fill = isi(ABU)
    b.alignment = TENGAH
    b.border = GARIS
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:F{F_AKHIR}"

# ===================================================================================================== lembar: Data Uji
ws = wb.create_sheet("Data Uji")
judul(ws, "Data uji", "Contoh isian untuk uji batas dan alat bantu hitung. Sel biru = angka yang boleh diubah (masukan); sel abu-abu = hasil rumus.", 7)
lebar_kolom(ws, [28, 22, 22, 22, 18, 22, 40])
r = 4
ws.cell(row=r, column=1, value="A. Kalkulator ambang lulus pelajaran (untuk TC-NL-04)").font = f(True, size=11, color=NAVY)
ws.cell(row=r + 1, column=1, value="Jumlah soal pelajaran").font = f(True)
masuk = ws.cell(row=r + 1, column=2, value=19)
masuk.font = f(True, color="0000FF")
masuk.fill = isi(KUNING)
masuk.border = GARIS
masuk.alignment = TENGAH
ws.cell(row=r + 1, column=3, value="Ubah sesuai pelajaran yang diuji (tabel jumlah soal di bawah). Lulus bila akurasi 80% atau lebih; bintang 80-89 = 1, 90-99 = 2, 100 = 3; XP pertama kali lulus 10/20/30 (checkpoint 50).").font = f(italic=True, color="595959")
ws.merge_cells(start_row=r + 1, start_column=3, end_row=r + 1, end_column=7)
ws.cell(row=r + 1, column=3).alignment = RATA_KIRI
ws.row_dimensions[r + 1].height = 30
kepala(ws, r + 3, ["Jumlah jawaban salah", "Total jawaban (soal + salah)", "Akurasi (%)", "Hasil", "Bintang", "XP pelajaran biasa"], 32)
for s in range(0, 9):
    rr = r + 4 + s
    ws.cell(row=rr, column=1, value=s)
    ws.cell(row=rr, column=2, value=f"=$B${r + 1}+A{rr}")
    ws.cell(row=rr, column=3, value=f"=ROUND($B${r + 1}/B{rr}*100,0)")
    ws.cell(row=rr, column=4, value=f'=IF(C{rr}>=80,"Lulus","Belum Lulus")')
    ws.cell(row=rr, column=5, value=f"=IF(C{rr}<80,0,IF(C{rr}=100,3,IF(C{rr}>=90,2,1)))")
    ws.cell(row=rr, column=6, value=f"=CHOOSE(E{rr}+1,0,10,20,30)")
    for kol in range(1, 7):
        c = ws.cell(row=rr, column=kol)
        c.border = GARIS
        c.font = f(bold=(kol == 1))
        c.alignment = TENGAH
        if kol > 1:
            c.fill = isi(ABU)
r = r + 4 + 9 + 1
ws.cell(row=r, column=1, value="B. Jumlah soal tiap pelajaran (data saat kit dibuat)").font = f(True, size=11, color=NAVY)
kepala(ws, r + 1, ["Pelajaran (ID)", "Jumlah soal", "Soal pilih", "Soal susun"], 22)
SOAL = [("l1 Baris a, ka, sa", 19, 15, 4), ("l2 Baris ta, na, ha", 19, 15, 4), ("l3 Baris ma, ya, ra, wa, n", 20, 16, 4), ("l4 Dakuten: ga, za, da", 19, 15, 4),
        ("l5 Dakuten ba dan handakuten pa", 14, 10, 4), ("l6 Youon 1", 16, 12, 4), ("l7 Youon 2, sokuon, vokal panjang", 21, 21, 0),
        ("k3_l1 sampai k3_l11 (Unit 3)", "21 sampai 29", "-", "-")]
for i, baris_soal in enumerate(SOAL):
    for kol, v in enumerate(baris_soal, start=1):
        c = ws.cell(row=r + 2 + i, column=kol, value=v)
        c.border = GARIS
        c.font = f(bold=(kol == 1))
        c.alignment = TENGAH if kol > 1 else RATA_KIRI
r = r + 2 + len(SOAL) + 1
ws.cell(row=r, column=1, value="C. Isian teks untuk uji batas").font = f(True, size=11, color=NAVY)
kepala(ws, r + 1, ["Jenis", "Nilai uji", "", "", "", "", "Yang diperiksa"], 22)
BATAS = [
    ("Nama", "(kosong)", "Ditolak dengan pesan jelas."),
    ("Nama", "     (spasi saja)", "Ditolak: \"Nama tidak boleh kosong.\"."),
    ("Nama", "A", "Diterima (1 karakter)."),
    ("Nama", "田中 さくら Ani", "Huruf Jepang dan Latin tampil benar di semua halaman."),
    ("Nama", "Ani 😀 Sari", "Emoji tidak merusak tampilan atau tersimpan sebagai tanda tanya."),
    ("Nama", "<script>alert(1)</script>", "Tampil sebagai teks, tidak dijalankan."),
    ("Nama", "' OR 1=1 --", "Tersimpan sebagai teks biasa."),
    ("Nama", "(300 karakter: Bu diulang 150 kali)", "Tidak merusak tata letak; server saat ini tidak membatasi panjang."),
    ("Email", "nama@contoh.id", "Valid."),
    ("Email", "NAMA@Contoh.ID", "Dianggap sama dengan huruf kecil."),
    ("Email", "nama@", "Ditolak oleh browser/server."),
    ("Email", "bukan-email", "Ditolak oleh browser/server."),
    ("Email", "nama+uji@contoh.id", "Diterima (tanda + sah)."),
    ("Password baru", "abc123", "Ditolak: kurang dari 8 karakter."),
    ("Password baru", "abcdefgh", "Ditolak: tidak ada angka."),
    ("Password baru", "12345678", "Ditolak: tidak ada huruf."),
    ("Password baru", "Murid2026", "Diterima."),
    ("Pencarian Kamus", "oishii | Enak | おいしい | お名前 | ENAK", "Menemukan kata yang sama; tidak peka huruf besar-kecil."),
    ("Pencarian Kamus", "zzzzbukankata | <b>x</b> | 600 huruf a", "Pesan Tidak ada hasil; tidak ada galat."),
    ("Kalimat uji AI (benar)", "はじめまして。わたしはアニです。", "Dibalas wajar."),
    ("Kalimat uji AI (salah tata bahasa)", "わたしは学生がです。", "Dikoreksi lewat balasan."),
    ("Kalimat uji AI (bukan Jepang)", "Halo, apa kabar?", "Tetap dalam peran, mengajak memakai bahasa Jepang."),
    ("Kalimat uji AI (500+ karakter)", "あ diulang 600 kali", "Kolom menahan di 500 karakter."),
    ("Skenario Perkenalan Diri (jawaban benar)", "Baris 1: はじめまして。よろしくお願いします。 | Baris 2: わたしの名前はアニです。", "Dipakai untuk TC-PC-07."),
]
for i, (jenis, nilai, periksa) in enumerate(BATAS):
    rr = r + 2 + i
    ws.cell(row=rr, column=1, value=jenis).font = f(True)
    ws.cell(row=rr, column=2, value=nilai).font = f()
    ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=6)
    ws.cell(row=rr, column=7, value=periksa).font = f(color="595959")
    for kol in (1, 2, 7):
        ws.cell(row=rr, column=kol).alignment = RATA_KIRI
    for kol in range(1, 8):
        ws.cell(row=rr, column=kol).border = GARIS
    ws.row_dimensions[rr].height = tinggi_baris([(nilai, 100), (periksa, 40)], minimum=16)

# ================================================================================================ lembar: Temuan Awal
ws = wb.create_sheet("Temuan Awal")
judul(ws, "Temuan awal saat menyusun buku ini", "Dari membaca kode dan menjalankan aplikasi di lingkungan pengembangan (bukan di server penguji). Ini bahan keputusan pemilik produk, bukan bug yang sudah dikonfirmasi di server Anda: "
      "konfirmasi dulu lewat kasus uji terkait.", 7)
lebar_kolom(ws, [5, 16, 46, 44, 40, 16, 40])
kepala(ws, 4, ["No", "Area", "Temuan", "Dampak", "Kasus uji terkait", "Tingkat", "Saran atau keputusan yang dibutuhkan"], 30)
TEMUAN = [
    ("Akun", "Tombol Keluar hanya menghapus token di perangkat. API tidak punya endpoint logout; token pembaruan lama tetap bisa dipakai sampai 14 hari atau sampai password diganti.",
     "Token yang sempat tercuri tidak bisa dicabut lewat Keluar.", "TC-AK-24, TC-AK-21", "Sedang", "Putuskan: tambah logout sisi server yang mencabut token pembaruan atau terima risikonya."),
    ("Admin", "Murid yang dinonaktifkan admin tidak bisa masuk lagi, tetapi sesi yang sedang berjalan tetap bisa dipakai (GET /me dan refresh tetap berhasil) sampai token pembaruan habis (14 hari).",
     "Akun yang seharusnya sudah dicabut aksesnya masih bisa belajar dan mendapat XP. PRD menulis \"murid tidak bisa login\".", "TC-AD-27, TC-AK-23", "Sedang",
     "Putuskan: cabut sesi saat akun dinonaktifkan. Skrip uji sudah memuat kasus yang diharapkan gagal sampai ini diputuskan."),
    ("Percakapan", "XP skenario Mode Tes (20) diberikan pada percobaan Tes PERTAMA walau percobaan itu gagal (Kesempatan Habis); lulus di percobaan berikutnya tidak memberi XP lagi.",
     "Murid bisa mendapat XP dengan sengaja gagal, dan yang lulus belakangan tidak diberi XP. Berbeda dari pelajaran yang hanya memberi XP saat lulus, dan dari catatan docs/PLAN.md (\"percobaan TEST pertama yang LULUS\").", "TC-PC-08", "Sedang",
     "Putuskan: XP hanya untuk Tes yang lulus pertama kali (disarankan), atau tetapkan sebagai aturan yang disengaja dan jelaskan di dokumen."),
    ("Email", "Email undangan, reset password, dan permintaan undangan ulang belum terkirim; tautannya hanya dicatat di log API (\"[STUB]\").",
     "Murid dan admin tidak menerima email apa pun. Sebelum dipakai murid sungguhan, tautan harus disalin dari log dan dikirim manual.", "TC-AK-31, TC-AD-14, TC-AD-28", "Tinggi",
     "Putuskan penyedia email dan sambungkan, atau tetapkan prosedur manual tertulis. Sudah tercatat di dokumen fitur (AK-10)."),
    ("Akun", "Layar \"Cek Emailmu\" pada Lupa password menulis \"kami sudah mengirim tautan\" padahal email belum terkirim (lihat temuan di atas).",
     "Murid menunggu email yang tidak akan datang.", "TC-AK-16", "Sedang", "Ubah teks sementara atau aktifkan email sebelum rilis."),
    ("Pelajaran", "Alamat pelajaran yang masih terkunci tetap bisa dibuka dan dikerjakan; hanya pengiriman hasilnya yang ditolak server di akhir (\"Gagal mengirim hasil belajar.\"). Isi soal berikut kunci jawabannya juga terkirim ke browser (sengaja, agar umpan balik instan).",
     "Murid bisa membuang waktu mengerjakan pelajaran terkunci dan melihat jawaban lewat alat pengembang. Tidak merusak data.", "TC-BR-05", "Rendah", "Terima sebagai perilaku yang dirancang, atau tolak membuka layar pelajaran terkunci."),
    ("Profil", "Nama murid tidak dibatasi panjangnya di server.", "Nama sangat panjang bisa merusak tata letak di tempat lain (di uji 300 karakter, tampilan aman).", "TC-PF-07", "Rendah", "Tetapkan batas panjang (mis. 100 karakter)."),
    ("Pelajaran", "Bila pengiriman hasil pelajaran gagal (jaringan putus), jawaban yang sudah dikerjakan hilang dan murid harus mengulang pelajaran.", "Murid kehilangan usaha 3-5 menit.", "TC-PL-14", "Rendah", "Tawarkan tombol \"Coba kirim lagi\" yang mengirim ulang log jawaban."),
    ("Admin", "Label peran admin di menu (abu-abu muda) kurang kontras terhadap latar putih; judul kolom tabel admin tidak bertanda scope sehingga dibaca sebagai sel biasa oleh pembaca layar.",
     "Aksesibilitas aplikasi admin.", "TC-AX-03, TC-AX-09", "Rendah", "Perbaiki warna teks dan tambahkan scope=\"col\" pada judul kolom."),
    ("Server", "Konfigurasi Caddy memberi header X-Content-Type-Options, X-Frame-Options, dan Referrer-Policy pada situs web, tetapi tidak memberi Content-Security-Policy maupun Strict-Transport-Security untuk situs web (API sudah memberi lewat pustaka keamanannya). Token masuk disimpan di penyimpanan browser.",
     "Tanpa CSP, skrip yang tersusupkan lebih leluasa mencuri token. Ini rekomendasi dari membaca konfigurasi, belum dibuktikan di server sungguhan.", "TC-NF-02", "Rendah", "Tambahkan HSTS dan CSP di Caddyfile setelah menguji dampaknya."),
    ("Konten", "Dua romaji di Unit 3 diduga salah ketik: \"paswādo\" (seharusnya pasuwādo) di kalimat k3s_2_4_2 dan \"raigetu\" (seharusnya raigetsu) di kalimat k3s_5_5_2. Ditemukan oleh pemeriksaan otomatis npm run konten.",
     "Murid mempelajari romaji yang salah.", "TC-KT-02", "Rendah", "Perbaiki di dokumen Word sumber lalu impor ulang; tinjau juga 6 kalimat yang sudah diduga keliru."),
    ("Konten", "Data uji belum punya audio sama sekali (367 butir kosakata dan kalimat) dan Kamus baru 169 kata (target produk 1.500).",
     "Fitur suara di pelajaran, Kamus, dan Flashcard belum bisa diuji; pencarian Kamus sering kosong.", "TC-KT-07, TC-KM-09, TC-KM-10", "Info", "Jalankan tts:seed di lingkungan uji bila ingin menguji suara."),
    ("Server", "Batas permintaan (100 per menit) dan batas permintaan undangan ulang (3 per jam) dihitung per alamat IP.",
     "Penguji di satu jaringan kantor berbagi batas; uji otomatis dan uji beban dari satu IP bisa terkena 429.", "TC-NF-05, TC-AK-05, TC-NF-11", "Info", "Jadwalkan uji otomatis tidak bersamaan; untuk uji beban naikkan batas di lingkungan uji."),
    ("Server", "Alamat di tautan undangan dan reset password diambil dari pengaturan CORS_ORIGIN_STUDENT di API (di Docker dibentuk dari PUBLIC_STUDENT_URL).", "Bila pengaturan ini masih localhost, tautan yang dikirim tidak bisa dibuka murid.", "TC-AD-30", "Info", "Periksa pengaturan ini di setiap deploy."),
]
for i, (area, temuan, dampak, kasus, tingkat, saran) in enumerate(TEMUAN, start=1):
    r = 4 + i
    for kol, v in enumerate([i, area, temuan, dampak, kasus, tingkat, saran], start=1):
        c = ws.cell(row=r, column=kol, value=v)
        c.font = f(bold=(kol == 1))
        c.alignment = TENGAH_ATAS if kol in (1, 6) else RATA_KIRI
        c.border = GARIS
    ws.row_dimensions[r].height = tinggi_baris([(temuan, 46), (dampak, 44), (kasus, 40), (saran, 40)])
warna_status(ws, f"F5:F{4 + len(TEMUAN)}", {"Tinggi": ("F8CBAD", "833C0B"), "Sedang": (JINGGA, "7F6000"), "Rendah": (HIJAU, "006100"), "Info": (BIRU_MUDA, "1F3864")})
ws.freeze_panes = "A5"

# ------------------------------------------------------------------------------------------------------- urutan lembar
urutan = ["Petunjuk", "Lingkungan", "Kasus Uji", "Ringkasan", "Log Bug", "Perangkat", "QC Rilis", "Cakupan Fitur", "Data Uji", "Temuan Awal"]
wb._sheets = [wb[n] for n in urutan]
for lembar in wb.worksheets:
    lembar.sheet_properties.tabColor = {"Petunjuk": NAVY, "Kasus Uji": "2E75B6", "Ringkasan": "548235", "Log Bug": "C00000"}.get(lembar.title, "A6A6A6")
for lembar in wb.worksheets:
    # Cetak: landscape, lebar muat satu halaman, tinggi bebas.
    lembar.page_setup.orientation = "landscape"
    lembar.page_setup.paperSize = lembar.PAPERSIZE_A4
    lembar.page_setup.fitToWidth = 1
    lembar.page_setup.fitToHeight = 0
    lembar.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    lembar.print_options.horizontalCentered = False
    lembar.page_margins.left = lembar.page_margins.right = 0.4
wb.active = 0
wb.properties.title = "Buku Kasus Uji dan QC: Aplikasi Belajar Bahasa Jepang"
wb.properties.creator = "Tim pengembang"
wb.save(KELUARAN)

# ------------------------------------------------------------- QC Rilis dalam bentuk teks (untuk dicetak / ditempel di tiket)
MD = AKAR / "checklists" / "QC-rilis.md"
MD.parent.mkdir(exist_ok=True)
baris_md = [
    "# Daftar periksa QC sebelum rilis",
    "",
    "> Berkas ini dibuat otomatis oleh `manual/build_workbook.py` dari daftar yang sama dengan lembar **QC Rilis** di `manual/Buku-Kasus-Uji-QA.xlsx`. Jangan diedit tangan: ubah daftarnya di skrip lalu jalankan ulang.",
    "> Item **Wajib** harus lulus (atau tidak berlaku dengan alasan tertulis) sebelum rilis. Hitungan otomatisnya ada di lembar Ringkasan buku Excel.",
    "",
    "Rilis / versi: ____________________   Tanggal: ____________   Penguji utama: ____________________   Penyetuju: ____________________",
    "",
]
kelompok_sebelumnya = None
nomor = 0
for kel, periksa, wajib, rujuk in QC:
    nomor += 1
    if kel != kelompok_sebelumnya:
        baris_md += ["", f"## {kel}", ""]
        kelompok_sebelumnya = kel
    tanda = "**Wajib**" if wajib == "Ya" else "pilihan"
    acuan = f" _(rujukan: {rujuk})_" if rujuk else ""
    baris_md.append(f"- [ ] {nomor}. {periksa} ({tanda}){acuan}")
baris_md += ["", "## Keputusan akhir", "", "- [ ] Layak rilis", "- [ ] Layak dengan catatan: ______________________________________", "- [ ] Tidak layak: ______________________________________", ""]
MD.write_text("\n".join(baris_md), encoding="utf-8")

otomatis = sum(1 for k in DAFTAR if k.id in SKRIP)
print(f"Tertulis {KELUARAN.name}: {len(DAFTAR)} kasus ({otomatis} punya skrip otomatis, {len(DAFTAR) - otomatis} manual), {len(fitur.FITUR)} fitur tercakup.")
